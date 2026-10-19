import io
import calendar
import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Alignment, Font, Border, Side

def crear_excel(df_res, mes=11, ano=2026):
    """
    Genera el archivo Excel final con formato mejorado:
    - Conserva el orden original de la plantilla.
    - Encabezados de columnas con fecha y nombre de día (ej. '1 Dom', '2 Lun').
    - Distintivo de color especial para los fines de semana (Sáb/Dom).
    - Paleta corporativa de colores para celdas de turnos.
    """
    salida = io.BytesIO()
    
    dias_espanol = ['Lun', 'Mar', 'Mié', 'Jue', 'Vie', 'Sáb', 'Dom']
    num_dias = calendar.monthrange(ano, mes)[1]
    primer_dia_semana, _ = calendar.monthrange(ano, mes)
    
    # Crear los encabezados formateados y mapear fines de semana
    renombrados = {}
    es_fin_de_semana = {}
    
    for d in range(1, num_dias + 1):
        dia_idx = (primer_dia_semana + d - 1) % 7
        nombre_dia = dias_espanol[dia_idx]
        col_name = f"{d} {nombre_dia}"
        renombrados[d] = col_name
        # Sábado (idx 5) y Domingo (idx 6)
        es_fin_de_semana[col_name] = (dia_idx in [5, 6])

    df_export = df_res.copy()
    df_export = df_export.rename(columns=renombrados)
    
    colores_turnos = {
        "C": "F4D35E",  # Amarillo (Corrido)
        "N": "5DADE2",  # Celeste (Noche)
        "P": "F5B7B1",  # Rosado (Post-turno)
        "D": "D5D8DC",  # Gris (Descanso)
        "V": "BB8FCE",  # Morado (Vacaciones)
        "L": "82E0AA",  # Verde (Libre)
    }

    fill_header_norm = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    fill_header_weekend = PatternFill(start_color="1B4F72", end_color="1B4F72", fill_type="solid") # Azul destacado para fines de semana
    font_header = Font(color="FFFFFF", bold=True, name="Calibri", size=10)
    font_dark = Font(color="202020", name="Calibri", size=11)
    align_center = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='D0D3D4'),
        right=Side(style='thin', color='D0D3D4'),
        top=Side(style='thin', color='D0D3D4'),
        bottom=Side(style='thin', color='D0D3D4')
    )

    with pd.ExcelWriter(salida, engine="openpyxl") as writer:
        df_export.to_excel(writer, sheet_name="Turnos", index=False)
        
        # Pestaña complementaria de diagnóstico
        diagnostico = df_res.attrs.get("diagnostico", {"estado": df_res.attrs.get("status", "OK")})
        pd.DataFrame([diagnostico]).to_excel(writer, sheet_name="Diagnostico", index=False)

        # Aplicar formato a la hoja Turnos
        wb = writer.book
        ws = wb["Turnos"]
        
        # 1. Formato de Encabezados (Fila 1)
        for col_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.alignment = align_center
            cell.font = font_header
            col_val = str(cell.value)
            
            if col_val in es_fin_de_semana and es_fin_de_semana[col_val]:
                cell.fill = fill_header_weekend
            else:
                cell.fill = fill_header_norm

        # 2. Formato y Colores de las Celdas de Turnos
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
            for cell in row:
                cell.alignment = align_center
                cell.font = font_dark
                cell.border = thin_border
                val = str(cell.value).strip().upper()
                
                if val in colores_turnos:
                    cell.fill = PatternFill(start_color=colores_turnos[val], end_color=colores_turnos[val], fill_type="solid")

    salida.seek(0)
    return salida.getvalue(), f"Turnos_{mes:02d}_{ano}.xlsx"
